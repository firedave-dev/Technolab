# -*- coding: utf-8 -*-
"""
Boite a outils Word : mise en page, styles, champs automatiques.

Les regles appliquees ici viennent du « Guide de redaction du Projet Tutore de
fin de cycle » de TechnoLAB-ISTA : Times New Roman, 16 pour les titres, 14 pour
les sous-titres, 12 pour le corps, interligne 1,5, texte justifie, marges de
2,54 cm, pagination, un chapitre par nouvelle page.

CHAMPS AUTOMATIQUES PLUTOT QUE TEXTE FIGE : la table des matieres, la liste des
figures, celle des tableaux et les numeros de figure sont inseres comme des
CHAMPS Word. Consequence pratique : apres avoir ajoute un paragraphe ou une
image, il suffit de selectionner tout le document (Ctrl+A) et d'appuyer sur F9
pour que numeros de page et numerotation se recalculent seuls. Un sommaire tape
a la main serait faux des la premiere correction.
"""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

POLICE = 'Times New Roman'
CORPS = 12
SOUS_TITRE = 14
TITRE = 16

# Mettre RGBColor(0, 0, 0) pour un document entierement noir.
COULEUR_TITRES = RGBColor(0x2E, 0x44, 0x74)
GRIS_LEGENDE = RGBColor(0x44, 0x44, 0x44)
ROUGE = RGBColor(0xB9, 0x1C, 0x1C)

# Largeur utile d'une page A4 avec des marges de 2,54 cm.
LARGEUR_UTILE = Cm(21.0 - 2 * 2.54)
HAUTEUR_MAX_FIGURE = Cm(19.0)


# ----------------------------------------------------------------- bas niveau
def _police(run, taille=CORPS, gras=False, italique=False, couleur=None):
    run.font.name = POLICE
    run.font.size = Pt(taille)
    run.font.bold = gras
    run.font.italic = italique
    if couleur is not None:
        run.font.color.rgb = couleur
    # Sans w:eastAsia, Word retombe sur une police de substitution pour
    # certains caracteres (guillemets, tirets cadratins) et la page devient
    # visiblement heterogene.
    run._element.rPr.rFonts.set(qn('w:eastAsia'), POLICE)
    return run


def champ(paragraphe, instruction, provisoire='…'):
    """
    Insere un champ Word (TOC, PAGE, SEQ…) qui se recalcule avec F9.

    `provisoire` est ce que Word affiche tant que le champ n'a pas ete mis a
    jour ; il disparait des la premiere actualisation.
    """
    run = paragraphe.add_run()
    debut = OxmlElement('w:fldChar')
    debut.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    separateur = OxmlElement('w:fldChar')
    separateur.set(qn('w:fldCharType'), 'separate')
    texte = OxmlElement('w:t')
    texte.text = provisoire
    fin = OxmlElement('w:fldChar')
    fin.set(qn('w:fldCharType'), 'end')
    for element in (debut, instr, separateur, texte, fin):
        run._r.append(element)
    return run


def actualiser_a_louverture(document):
    """Demande a Word de recalculer tous les champs a l'ouverture du fichier."""
    parametres = document.settings.element
    balise = OxmlElement('w:updateFields')
    balise.set(qn('w:val'), 'true')
    parametres.append(balise)


# ------------------------------------------------------------- mise en page
def preparer(document):
    """Applique la charte typographique du guide au style par defaut."""
    normal = document.styles['Normal']
    normal.font.name = POLICE
    normal.font.size = Pt(CORPS)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), POLICE)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)

    for section in document.sections:
        marges(section)

    # Les styles de titre servent a alimenter la table des matieres : le champ
    # TOC ne sait reperer que les paragraphes portant un style « Titre n ».
    for nom, taille in [('Heading 1', TITRE), ('Heading 2', SOUS_TITRE), ('Heading 3', CORPS)]:
        style = document.styles[nom]
        style.font.name = POLICE
        style.font.size = Pt(taille)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = COULEUR_TITRES
        style.element.rPr.rFonts.set(qn('w:eastAsia'), POLICE)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True
    return document


def marges(section, cm=2.54):
    section.top_margin = Cm(cm)
    section.bottom_margin = Cm(cm)
    section.left_margin = Cm(cm)
    section.right_margin = Cm(cm)


def numeroter(section, format_='decimal', depart=None):
    """
    Fixe le format de pagination d'une section.

    Le guide veut des pages numerotees ; l'usage academique numerote les pages
    liminaires en chiffres romains et le corps en chiffres arabes repartant de 1.
    """
    sectPr = section._sectPr
    balise = sectPr.find(qn('w:pgNumType'))
    if balise is None:
        balise = OxmlElement('w:pgNumType')
        sectPr.append(balise)
    balise.set(qn('w:fmt'), format_)
    if depart is not None:
        balise.set(qn('w:start'), str(depart))


def pied_de_page(section, avec_numero=True):
    """Numero de page centre en pied de page."""
    pied = section.footer
    pied.is_linked_to_previous = False
    paragraphe = pied.paragraphs[0]
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphe.paragraph_format.line_spacing = 1.0
    for run in list(paragraphe.runs):
        run._element.getparent().remove(run._element)
    if avec_numero:
        _police(champ(paragraphe, 'PAGE', '1'), taille=11)


def nouvelle_section(document, format_pagination='decimal', depart=None, numerotee=True):
    """Ouvre une section sur une nouvelle page, avec sa propre pagination."""
    section = document.add_section(WD_SECTION.NEW_PAGE)
    marges(section)
    section.header.is_linked_to_previous = False
    for paragraphe in section.header.paragraphs:
        for run in list(paragraphe.runs):
            run._element.getparent().remove(run._element)
    numeroter(section, format_pagination, depart)
    pied_de_page(section, numerotee)
    return section


def saut_de_page(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------- contenus
def para(document, texte='', taille=CORPS, gras=False, italique=False,
         alignement=WD_ALIGN_PARAGRAPH.JUSTIFY, couleur=None, espace_avant=0,
         interligne=1.5):
    paragraphe = document.add_paragraph()
    paragraphe.alignment = alignement
    paragraphe.paragraph_format.space_before = Pt(espace_avant)
    paragraphe.paragraph_format.line_spacing = interligne
    if texte:
        _police(paragraphe.add_run(texte), taille, gras, italique, couleur)
    return paragraphe


def riche(document, morceaux, alignement=WD_ALIGN_PARAGRAPH.JUSTIFY, taille=CORPS):
    """
    Paragraphe melant du texte normal et du gras.

    `morceaux` est une suite de chaines ou de couples (texte, True) pour le gras.
    """
    paragraphe = document.add_paragraph()
    paragraphe.alignment = alignement
    paragraphe.paragraph_format.line_spacing = 1.5
    for morceau in morceaux:
        texte, gras = morceau if isinstance(morceau, tuple) else (morceau, False)
        _police(paragraphe.add_run(texte), taille, gras)
    return paragraphe


def titre1(document, texte, nouvelle_page=True):
    if nouvelle_page:
        saut_de_page(document)
    paragraphe = document.add_heading(level=1)
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _police(paragraphe.add_run(texte), TITRE, True, couleur=COULEUR_TITRES)
    return paragraphe


def titre2(document, texte):
    paragraphe = document.add_heading(level=2)
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _police(paragraphe.add_run(texte), SOUS_TITRE, True, couleur=COULEUR_TITRES)
    return paragraphe


def titre3(document, texte):
    paragraphe = document.add_heading(level=3)
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _police(paragraphe.add_run(texte), CORPS, True, couleur=COULEUR_TITRES)
    return paragraphe


def puces(document, elements, symbole='—'):
    for element in elements:
        paragraphe = document.add_paragraph()
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraphe.paragraph_format.left_indent = Cm(0.8)
        paragraphe.paragraph_format.first_line_indent = Cm(-0.5)
        paragraphe.paragraph_format.space_after = Pt(3)
        paragraphe.paragraph_format.line_spacing = 1.5
        if isinstance(element, tuple):
            tete, suite = element
            _police(paragraphe.add_run(f'{symbole} '), CORPS)
            _police(paragraphe.add_run(tete), CORPS, gras=True)
            _police(paragraphe.add_run(suite), CORPS)
        else:
            _police(paragraphe.add_run(f'{symbole} {element}'), CORPS)


# ----------------------------------------------------------------- figures
def figure(document, chemin, legende, largeur=None):
    """
    Image centree, legende NUMEROTEE AUTOMATIQUEMENT en dessous.

    Le guide impose la legende sous la figure (et au-dessus des tableaux). Le
    numero est un champ SEQ : inserer une figure au milieu du document renumerote
    toutes les suivantes sans intervention.
    """
    if not chemin.exists():
        return cadre_absent(document, legende, chemin.name)

    largeur = largeur or LARGEUR_UTILE
    with Image.open(chemin) as image:
        rapport_hw = image.height / image.width
    if largeur * rapport_hw > HAUTEUR_MAX_FIGURE:
        largeur = int(HAUTEUR_MAX_FIGURE / rapport_hw)

    paragraphe = document.add_paragraph()
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphe.paragraph_format.space_before = Pt(10)
    paragraphe.paragraph_format.space_after = Pt(2)
    paragraphe.paragraph_format.line_spacing = 1.0
    paragraphe.add_run().add_picture(str(chemin), width=largeur)
    return legende_figure(document, legende)


def legende_figure(document, legende):
    paragraphe = document.add_paragraph()
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphe.paragraph_format.space_after = Pt(12)
    paragraphe.paragraph_format.line_spacing = 1.0
    paragraphe.style = document.styles['Caption']
    _police(paragraphe.add_run('Figure '), 11, gras=True, couleur=GRIS_LEGENDE)
    _police(champ(paragraphe, r' SEQ Figure \* ARABIC ', '1'), 11, gras=True, couleur=GRIS_LEGENDE)
    _police(paragraphe.add_run(f' : {legende}'), 11, couleur=GRIS_LEGENDE)
    return paragraphe


def cadre_absent(document, legende, fichier):
    """
    Emplacement reserve a une capture d'ecran non encore fournie.

    Le rapport se genere malgre l'absence du fichier : mieux vaut un cadre
    visible, qu'on ne peut pas oublier, qu'une figure manquante silencieuse.
    """
    tableau = document.add_table(rows=1, cols=1)
    tableau.style = 'Table Grid'
    tableau.alignment = WD_TABLE_ALIGNMENT.CENTER
    cellule = tableau.cell(0, 0)
    cellule.width = LARGEUR_UTILE
    paragraphe = cellule.paragraphs[0]
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphe.paragraph_format.space_before = Pt(38)
    paragraphe.paragraph_format.space_after = Pt(38)
    _police(paragraphe.add_run(f'Capture à insérer\n{fichier}'), 11,
            italique=True, couleur=ROUGE)
    return legende_figure(document, legende)


# ----------------------------------------------------------------- tableaux
def tableau(document, legende, entetes, lignes, source='Les auteurs',
            largeurs=None):
    """
    Tableau a legende AU-DESSUS, conformement au guide, et source en dessous.
    """
    paragraphe = document.add_paragraph()
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraphe.paragraph_format.space_before = Pt(10)
    paragraphe.paragraph_format.space_after = Pt(3)
    paragraphe.paragraph_format.line_spacing = 1.0
    paragraphe.style = document.styles['Caption']
    _police(paragraphe.add_run('Tableau '), 11, gras=True, couleur=GRIS_LEGENDE)
    _police(champ(paragraphe, r' SEQ Tableau \* ARABIC ', '1'), 11, gras=True, couleur=GRIS_LEGENDE)
    _police(paragraphe.add_run(f' : {legende}'), 11, couleur=GRIS_LEGENDE)

    table = document.add_table(rows=1, cols=len(entetes))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cellule, texte in zip(table.rows[0].cells, entetes):
        cellule.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cellule.paragraphs[0].paragraph_format.line_spacing = 1.0
        cellule.paragraphs[0].paragraph_format.space_after = Pt(2)
        _police(cellule.paragraphs[0].add_run(texte), 10.5, gras=True)
        _ombrer(cellule, 'DCE3EE')

    for ligne in lignes:
        cellules = table.add_row().cells
        for index, (cellule, texte) in enumerate(zip(cellules, ligne)):
            cellule.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index == 0 and len(str(texte)) <= 6
                else WD_ALIGN_PARAGRAPH.LEFT)
            cellule.paragraphs[0].paragraph_format.line_spacing = 1.0
            cellule.paragraphs[0].paragraph_format.space_after = Pt(2)
            _police(cellule.paragraphs[0].add_run(str(texte)), 10.5)

    if largeurs:
        for ligne in table.rows:
            for cellule, largeur in zip(ligne.cells, largeurs):
                cellule.width = Cm(largeur)

    if source:
        pied = document.add_paragraph()
        pied.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pied.paragraph_format.space_after = Pt(12)
        pied.paragraph_format.line_spacing = 1.0
        _police(pied.add_run(f'Source : {source}'), 10, italique=True, couleur=GRIS_LEGENDE)
    return table


def _ombrer(cellule, couleur_hex):
    ombre = OxmlElement('w:shd')
    ombre.set(qn('w:val'), 'clear')
    ombre.set(qn('w:fill'), couleur_hex)
    cellule._tc.get_or_add_tcPr().append(ombre)
